import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# Colors
COLOR_PRIMARY = RGBColor(31, 78, 121)     # #1F4E79 - Dark Steel Blue
COLOR_SECONDARY = RGBColor(47, 85, 151)   # #2F5597 - Classic Blue
COLOR_BODY = RGBColor(51, 51, 51)         # #333333 - Charcoal
COLOR_TEXT_MUTED = RGBColor(128, 128, 128)
HEX_PRIMARY = "1F4E79"
HEX_LIGHT_GRAY = "F2F4F7"
HEX_YELLOW_BG = "FCF3CF"                 # For screenshots
COLOR_SCREENSHOT_BG = "FCF3CF"

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_page_number(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> PAGE </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def add_total_pages(run):
    fldChar1 = parse_xml(r'<w:fldChar %s w:fldCharType="begin"/>' % nsdecls('w'))
    instrText = parse_xml(r'<w:instrText %s xml:space="preserve"> NUMPAGES </w:instrText>' % nsdecls('w'))
    fldChar2 = parse_xml(r'<w:fldChar %s w:fldCharType="separate"/>' % nsdecls('w'))
    fldChar3 = parse_xml(r'<w:fldChar %s w:fldCharType="end"/>' % nsdecls('w'))
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def setup_header_footer(section, doc_title):
    section.different_first_page_header_footer = True
    
    # Header Setup
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(doc_title)
    hrun.font.name = 'Segoe UI'
    hrun.font.size = Pt(8.5)
    hrun.font.color.rgb = COLOR_TEXT_MUTED
    
    # Footer Setup
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.text = "" # Clear default
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), 2) # Right align tab
    
    r_left = fp.add_run("Dataverse Label Translator  |  E2E User Scenario")
    r_left.font.name = 'Segoe UI'
    r_left.font.size = Pt(8.5)
    r_left.font.color.rgb = COLOR_TEXT_MUTED
    
    fp.add_run("\tPage ")
    r_page = fp.add_run()
    r_page.font.name = 'Segoe UI'
    r_page.font.size = Pt(8.5)
    r_page.font.color.rgb = COLOR_TEXT_MUTED
    add_page_number(r_page)
    
    fp.add_run(" of ")
    r_total = fp.add_run()
    r_total.font.name = 'Segoe UI'
    r_total.font.size = Pt(8.5)
    r_total.font.color.rgb = COLOR_TEXT_MUTED
    add_total_pages(r_total)

def add_heading_1(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(18)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(16)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    
    pPr = p._p.get_or_add_pPr()
    pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="4" w:color="{HEX_PRIMARY}"/></w:pBdr>')
    pPr.append(pbdr)
    return p

def add_heading_2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(13)
    run.bold = True
    run.font.color.rgb = COLOR_SECONDARY
    return p

def add_heading_3(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.keep_with_next = True
    
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(11)
    run.bold = True
    run.font.color.rgb = COLOR_PRIMARY
    return p

def add_body_text(doc, text, bold_prefix=None, space_after=6, bullet=False):
    p = doc.add_paragraph(style='List Bullet' if bullet else 'Normal')
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r_prefix = p.add_run(bold_prefix)
        r_prefix.font.name = 'Segoe UI'
        r_prefix.font.size = Pt(10.5)
        r_prefix.bold = True
        r_prefix.font.color.rgb = COLOR_BODY
        
    run = p.add_run(text)
    run.font.name = 'Segoe UI'
    run.font.size = Pt(10.5)
    run.font.color.rgb = COLOR_BODY
    return p

def create_test_table(doc, test_num, persona, problem, pain_points, user_goals, business_goals, steps, success_metrics):
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(12)
    p_title.paragraph_format.space_after = Pt(12)
    p_title.paragraph_format.keep_with_next = True
    r_title = p_title.add_run(f"Test Case: {test_num} - {persona}")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(13)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY

    # Table with 8 rows and 2 columns
    table = doc.add_table(rows=8, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Set widths (approx 1.8 inches and 4.7 inches)
    widths = [Inches(1.8), Inches(4.7)]
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width
            
    def style_cell(cell, is_header=False, is_left_col=False):
        set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
        if is_left_col:
            set_cell_background(cell, HEX_LIGHT_GRAY)
        elif is_header:
            set_cell_background(cell, HEX_PRIMARY)
            
    cell_l0 = table.cell(0, 0)
    cell_r0 = table.cell(0, 1)
    cell_l0.text = "Test Number"
    cell_r0.text = test_num
    style_cell(cell_l0, is_left_col=True)
    style_cell(cell_r0)
    
    for p in cell_l0.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.name = 'Segoe UI'
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_PRIMARY
    for p in cell_r0.paragraphs:
        for r in p.runs:
            r.bold = True
            r.font.name = 'Segoe UI'
            r.font.size = Pt(10)
            r.font.color.rgb = COLOR_PRIMARY
            
    row_data = [
        (1, "Primary Persona:", persona),
        (2, "Problem Statement:", problem),
        (3, "Pain Points:", pain_points),
        (4, "User Goals:", user_goals),
        (5, "Business Goals:", business_goals),
        (6, "Detailed Steps", steps),
        (7, "Success Metrics:", success_metrics)
    ]
    
    for r_idx, label, val in row_data:
        cell_l = table.cell(r_idx, 0)
        cell_r = table.cell(r_idx, 1)
        
        cell_l.text = label
        style_cell(cell_l, is_left_col=True)
        for p in cell_l.paragraphs:
            for r in p.runs:
                r.bold = True
                r.font.name = 'Segoe UI'
                r.font.size = Pt(9.5)
                r.font.color.rgb = COLOR_BODY
                
        style_cell(cell_r)
        
        if label == "Detailed Steps":
            p_default = cell_r.paragraphs[0]
            pPr = p_default._p.get_or_add_pPr()
            
            steps_list = val.split("\n")
            for idx_step, step in enumerate(steps_list):
                step_str = step.strip()
                if not step_str:
                    continue
                
                if "[📷 HÌNH ẢNH" in step_str or "[📷 SCREENSHOT" in step_str:
                    p_ss = cell_r.add_paragraph() if idx_step > 0 else p_default
                    p_ss.paragraph_format.space_before = Pt(6)
                    p_ss.paragraph_format.space_after = Pt(6)
                    p_ss.paragraph_format.left_indent = Inches(0.15)
                    p_ss.paragraph_format.right_indent = Inches(0.15)
                    
                    pPr_ss = p_ss._p.get_or_add_pPr()
                    pbdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="6" w:color="FFC000"/></w:pBdr>')
                    pPr_ss.append(pbdr)
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{HEX_YELLOW_BG}"/>')
                    pPr_ss.append(shd)
                    
                    r_ss = p_ss.add_run(step_str)
                    r_ss.font.name = 'Segoe UI'
                    r_ss.font.size = Pt(9)
                    r_ss.italic = True
                    r_ss.font.color.rgb = RGBColor(180, 100, 0)
                else:
                    p_step = cell_r.add_paragraph() if idx_step > 0 else p_default
                    p_step.paragraph_format.space_after = Pt(4)
                    p_step.paragraph_format.line_spacing = 1.15
                    
                    r_step = p_step.add_run(step_str)
                    r_step.font.name = 'Segoe UI'
                    r_step.font.size = Pt(9.5)
                    r_step.font.color.rgb = COLOR_BODY
                    
                    if "passed this test" in step_str.lower():
                        r_step.bold = True
                        r_step.font.color.rgb = COLOR_SECONDARY
        else:
            p_default = cell_r.paragraphs[0]
            lines = val.split("\n")
            for idx_line, line in enumerate(lines):
                line_str = line.strip()
                if not line_str:
                    continue
                p_line = cell_r.add_paragraph() if idx_line > 0 else p_default
                p_line.paragraph_format.space_after = Pt(4)
                p_line.paragraph_format.line_spacing = 1.15
                
                r_line = p_line.add_run(line_str)
                r_line.font.name = 'Segoe UI'
                r_line.font.size = Pt(9.5)
                r_line.font.color.rgb = COLOR_BODY

def build_e2e_document():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        setup_header_footer(section, "Dataverse Label Translator - E2E User Scenarios")
        
    # --- TITLE / COVER PAGE ---
    p_spacer = doc.add_paragraph()
    p_spacer.paragraph_format.space_before = Pt(80)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run("DATAVERSE LABEL TRANSLATOR")
    r_title.font.name = 'Segoe UI'
    r_title.font.size = Pt(28)
    r_title.bold = True
    r_title.font.color.rgb = COLOR_PRIMARY
    
    p_subtitle = doc.add_paragraph()
    p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_subtitle.paragraph_format.space_after = Pt(140)
    r_sub = p_subtitle.add_run("E2E User Scenario Testing Document for Microsoft AppSource")
    r_sub.font.name = 'Segoe UI'
    r_sub.font.size = Pt(14)
    r_sub.font.color.rgb = COLOR_TEXT_MUTED
    
    p_info = doc.add_paragraph()
    p_info.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_info.paragraph_format.space_after = Pt(6)
    r_info = p_info.add_run("Publisher: Phuoc Le\nVersion: 1.0.0.0\nPlatform: Microsoft Power Platform / Dataverse")
    r_info.font.name = 'Segoe UI'
    r_info.font.size = Pt(11)
    r_info.font.color.rgb = COLOR_BODY
    
    doc.add_page_break()

    # --- INTRODUCTORY PAGE (Assumptions & TOC) ---
    add_heading_1(doc, "Document Overview & Assumptions")
    
    add_heading_2(doc, "Introduction")
    add_body_text(doc, "This document outlines the End-to-End (E2E) User Scenario Testing cases required for validating the Dataverse Label Translator model-driven application during deployment on Microsoft AppSource. Each scenario defines target personas, pain points, detailed step-by-step procedures, and success criteria.")
    
    add_heading_2(doc, "Assumptions & Prerequisites")
    add_body_text(doc, "Test with clean Microsoft Dataverse environment installed.", bullet=True)
    add_body_text(doc, "Solution DataverseLabelTranslator is successfully imported.", bullet=True)
    add_body_text(doc, "At least two user interface languages (e.g. English base and Vietnamese target) are enabled and provisioned in the environment.", bullet=True)
    add_body_text(doc, "The tester has access to one System Administrator/Customizer user account and one regular user account for security check testing.", bullet=True)
    
    add_heading_2(doc, "Test Cases Overview")
    add_body_text(doc, "TEST00: Regular User / Non-Admin User - Verify Security Role Requirements (Fail/Pass)", bullet=True)
    add_body_text(doc, "TEST01: Customizer/System Administrator - Translate Entity Attribute Labels (Manual Inline Translation)", bullet=True)
    add_body_text(doc, "TEST02: Customizer/System Administrator - All-In-One Translation for Custom Entity", bullet=True)
    add_body_text(doc, "TEST03: Customizer/System Administrator - AI-Assisted Bulk Translation", bullet=True)
    add_body_text(doc, "TEST04: Customizer/System Administrator - Manage and Apply Translation Dictionary (Glossary)", bullet=True)
    add_body_text(doc, "TEST05: Customizer/System Administrator - Translate Form Labels with Automatic User Language Switching", bullet=True)
    
    doc.add_page_break()
    
    # TEST00 (Page 3)
    persona_0 = "Regular User / Non-Admin User - Verify Security Role Requirements"
    problem_0 = (
        "+ You want to verify that metadata translation and customization are protected security-wise.\n"
        "+ Only users with either System Administrator or System Customizer security roles can load and modify system labels.\n"
        "+ Regular users without either of these roles must be blocked (fail), whereas users with either role must pass."
    )
    pain_points_0 = "+ Letting regular users modify system metadata or localization keys poses a massive security risk and can lead to data schema inconsistencies."
    user_goals_0 = "+ Ensure security validation prevents unauthorized users from loading and modifying translations, while letting customizers work seamlessly."
    business_goals_0 = "+ Prevent unauthorized system customizations and enforce strict environment governance policies."
    steps_0 = (
        "+ Log in to the Dataverse environment using a non-admin account (e.g. Salesperson, basic member) which does NOT have System Administrator or System Customizer roles.\n"
        "+ Launch the Dataverse Label Translator model-driven app.\n"
        "+ Try to select an Entity and click 'Load'.\n"
        "+ Verify that the application shows an error dialog (Access Denied / Insufficient Privileges) or blocks the loading action.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp ảnh màn hình thông báo lỗi phân quyền hoặc giao diện bị khóa khi đăng nhập bằng tài khoản không có quyền Admin/Customizer rồi paste vào đây.]\n"
        "+ This indicates the security restriction works (Test status: FAIL for unauthorized access - expected behavior).\n"
        "+ Now, log out and log back in with an account that has either the 'System Administrator' or 'System Customizer' security role.\n"
        "+ Re-open the app, select an Entity, and click 'Load'.\n"
        "+ Verify that the metadata loads successfully and the grid displays editable rows.\n"
        "+ You passed this test (Test status: PASS for authorized security role)."
    )
    metrics_0 = "+ Unauthorized users are successfully blocked from performing translation actions, while users with System Administrator or System Customizer roles are granted full access."
    
    create_test_table(doc, "TEST00", persona_0, problem_0, pain_points_0, user_goals_0, business_goals_0, steps_0, metrics_0)
    doc.add_page_break()
    
    # TEST01 (Page 4)
    persona_1 = "Customizer/System Administrator - Manual Attribute Translation"
    problem_1 = "+ You are a Customizer/System Administrator and\n+ Dataverse Label Translator is installed and\n+ You need to translate or customize field/attribute display names to a target language (e.g. Vietnamese)."
    pain_points_1 = "+ Exporting translation zip files, manually modifying XML translation strings, and re-importing is complex, slow, and highly prone to schema errors."
    user_goals_1 = "+ Directly translate and preview field labels in a clean interactive grid and save immediately."
    business_goals_1 = "+ Fast metadata localization, reducing time-to-market for multi-language deployments."
    steps_1 = (
        "+ Open the Dataverse Label Translator model-driven app.\n"
        "+ Select the Solution (e.g. Default Solution) from the Solution filter.\n"
        "+ Select the Entity (e.g. Account) from the dropdown.\n"
        "+ Select Type as '1. Attributes' and click the Load button.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy mở app, load Attributes của Account và chụp bảng grid chứa dữ liệu rồi paste vào đây.]\n"
        "+ Double-click in the target language column (e.g. Vietnamese) for a few attribute rows (e.g. Name -> Tên Công Ty, Main Phone -> Điện Thoại Chính).\n"
        "+ Observe the cell highlight (dirty indicator).\n"
        "+ Click the Save button in the top toolbar. Wait for the success dialog.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp ảnh màn hình thông báo Save thành công rồi paste vào đây.]\n"
        "+ You passed this test."
    )
    metrics_1 = "+ The customized field labels are successfully saved in Dataverse and displayed to end-users when switching their UI language."
    
    create_test_table(doc, "TEST01", persona_1, problem_1, pain_points_1, user_goals_1, business_goals_1, steps_1, metrics_1)
    doc.add_page_break()
    
    # TEST02 (Page 5)
    persona_2 = "Customizer/System Administrator - All-In-One Bulk Translation"
    problem_2 = "+ You are a Customizer/System Administrator and\n+ You have just created a new custom entity and\n+ You need to translate all related metadata (attributes, options, forms, views, etc.) without switching between individual type menus."
    pain_points_2 = "+ Switching back and forth between different components (Forms, Views, Fields) to load and save them separately is extremely tedious."
    user_goals_2 = "+ Load all entity-dependent labels into a single grid for fast bulk translations."
    business_goals_2 = "+ Accelerate localization of newly developed custom solutions."
    steps_2 = (
        "+ Select the Solution containing your custom entity.\n"
        "+ Select the newly created custom Entity.\n"
        "+ Select Type as '0. All-In-One'.\n"
        "+ Click the Load button to fetch all metadata.\n"
        "+ Verify that the grid displays rows representing different types (Attributes, Forms, Views, Option Sets, Relationships, etc.).\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp bảng Grid sau khi load chế độ All-In-One hiển thị nhiều loại component khác nhau rồi paste vào đây.]\n"
        "+ Enter translations for multiple components across different types.\n"
        "+ Click Save in the toolbar.\n"
        "+ Verify all component labels are saved and published in Dataverse.\n"
        "+ You passed this test."
    )
    metrics_2 = "+ All entity-related labels are loaded, edited, and saved in a single unified operation."
    
    create_test_table(doc, "TEST02", persona_2, problem_2, pain_points_2, user_goals_2, business_goals_2, steps_2, metrics_2)
    doc.add_page_break()
    
    # TEST03 (Page 6)
    persona_3 = "Customizer/System Administrator - AI-Assisted Bulk Translation"
    problem_3 = "+ You are a Customizer/System Administrator and\n+ You need to translate hundreds of metadata labels for an entity, and doing it manually is too slow."
    pain_points_3 = "+ Translating massive lists of fields manually requires translation services, copy-pasting, and is extremely slow."
    user_goals_3 = "+ Use advanced AI (Gemini, OpenAI, or Azure Foundry) to auto-translate empty cells in bulk, while ensuring technical terms and product names remain accurate."
    business_goals_3 = "+ Automate translation to cut down localization costs and timeline by 90%."
    steps_3 = (
        "+ Click the AI Settings button in the toolbar.\n"
        "+ Configure your AI Provider (e.g. Google Gemini or OpenAI) by pasting your API Key and selecting/typing the model.\n"
        "+ Click Save on the settings dialog.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp popup AI Settings sau khi điền khóa API rồi paste vào đây.]\n"
        "+ Select an Entity and Type (e.g. 1. Attributes) and click Load.\n"
        "+ Select the Source Language and Target Language in the toolbar.\n"
        "+ Click the Auto Translate button.\n"
        "+ In the popup, choose 'Missing values' and check the desired options.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp popup cấu hình Auto Translate rồi paste vào đây.]\n"
        "+ Click Translate and wait for the AI to populate the blank cells in green italic text.\n"
        "+ Review the proposals, modify any rows if needed, and click Save.\n"
        "+ You passed this test."
    )
    metrics_3 = "+ The tool successfully calls the AI provider, returns translation suggestions, populates the grid, and saves the values to Dataverse."
    
    create_test_table(doc, "TEST03", persona_3, problem_3, pain_points_3, user_goals_3, business_goals_3, steps_3, metrics_3)
    doc.add_page_break()
    
    # TEST04 (Page 7)
    persona_4 = "Customizer/System Administrator - Term Consistency via Glossary"
    problem_4 = "+ You are a Customizer/System Administrator and\n+ You want certain corporate terms, product names, or acronyms to be translated consistently, avoiding AI hallucinations or generic terms."
    pain_points_4 = "+ Manual and AI translations can be inconsistent across fields, forms, and views, causing user confusion."
    user_goals_4 = "+ Maintain a central glossary of terms that applies automatically to matching base language labels."
    business_goals_4 = "+ Maintain corporate branding and interface terminology consistency across the entire Dynamics 365 tenant."
    steps_4 = (
        "+ Click the Manage Dictionary button.\n"
        "+ In the popup dialog, enter a base language term (e.g. 'Account') and its corresponding target translation (e.g. 'Khách hàng' for Vietnamese).\n"
        "+ Click Save in the dictionary popup.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp popup Dictionary Editor hiển thị danh sách từ điển rồi paste vào đây.]\n"
        "+ Load the Attributes of the Account entity.\n"
        "+ Click the Apply Dictionary button in the toolbar.\n"
        "+ Select 'All Missing' or 'All Overwrite'.\n"
        "+ Verify that all grid cells containing the base word 'Account' are updated to the dictionary definition.\n"
        "+ Click Save to commit changes.\n"
        "+ You passed this test."
    )
    metrics_4 = "+ The dictionary entries are saved to Dataverse XML storage, and applying the dictionary updates matching cell values correctly."
    
    create_test_table(doc, "TEST04", persona_4, problem_4, pain_points_4, user_goals_4, business_goals_4, steps_4, metrics_4)
    doc.add_page_break()
    
    # TEST05 (Page 8)
    persona_5 = "Customizer/System Administrator - Multi-Language Form Label Translation"
    problem_5 = "+ You are a Customizer/System Administrator and\n+ You need to translate form labels (tabs, sections, fields) across multiple installed languages, but Dataverse API only returns labels for the currently active user language."
    pain_points_5 = "+ Changing the user language in personal settings, loading/saving forms, and repeating this for every single language is extremely slow and tedious."
    user_goals_5 = "+ Load all form labels for all languages and save them in one operation while the app automatically handles background language switching."
    business_goals_5 = "+ Enable effortless multi-language form design without manual settings navigation."
    steps_5 = (
        "+ Select the Entity (e.g. Contact).\n"
        "+ Select Type as '3. Forms' and click Load.\n"
        "+ Observe the loader status indicating that the application is temporarily changing the user language in the background to fetch labels for each language.\n"
        "+ [📷 HÌNH ẢNH HƯỚNG DẪN: Anh Phước hãy chụp thanh trạng thái/loader khi đang load Form để thể hiện việc đổi ngôn ngữ rồi paste vào đây.]\n"
        "+ Verify that the form tab/section/field labels are loaded for all columns.\n"
        "+ Translate several form labels (e.g., General Tab -> Thông tin chung, Summary -> Tóm tắt).\n"
        "+ Click Save.\n"
        "+ Observe the status bar showing the background language switching and publishing.\n"
        "+ Ensure the app restores your user interface back to your original language.\n"
        "+ You passed this test."
    )
    metrics_5 = "+ Form labels for all active languages are fetched, edited, successfully saved, and the administrator's UI language is restored."
    
    create_test_table(doc, "TEST05", persona_5, problem_5, pain_points_5, user_goals_5, business_goals_5, steps_5, metrics_5)
    
    # Save E2E Document
    version = "1.0.0.0"
    output_dir = rf"d:\github\DataverseLabelTranslator\release\{version}\appsource\Test"
    os.makedirs(output_dir, exist_ok=True)
    file_path = os.path.join(output_dir, f"E2E User Scenario.{version}.docx")
    doc.save(file_path)
    print(f"Document saved successfully at: {file_path}")

if __name__ == "__main__":
    build_e2e_document()
